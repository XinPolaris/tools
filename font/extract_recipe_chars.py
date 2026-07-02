# -*- coding: utf-8 -*-
"""从菜谱 Excel 的所有单元格中提取并去重汉字。"""

import argparse
import re
from pathlib import Path

from openpyxl import load_workbook


BASE_DIR = Path(__file__).resolve().parent
CHARS_DIR = BASE_DIR / "input" / "chars"
DEFAULT_INPUT = BASE_DIR / "assest" / "☆菜谱库信息汇总V2（含标签及步骤修改6.12）.xlsx"
DEFAULT_OUTPUT = CHARS_DIR / "recipe_chars.txt"

# CJK 基本区、扩展区及兼容汉字区。
HAN_PATTERN = re.compile(
    r"[\u3400-\u4DBF"
    r"\u4E00-\u9FFF"
    r"\uF900-\uFAFF"
    r"\U00020000-\U0002FA1F]"
)


def extract_unique_han_chars(xlsx_path: Path) -> str:
    """按首次出现顺序返回工作簿中的所有不重复汉字。"""
    workbook = load_workbook(xlsx_path, read_only=True, data_only=False)
    seen: set[str] = set()
    chars: list[str] = []

    try:
        for worksheet in workbook.worksheets:
            for row in worksheet.iter_rows(values_only=True):
                for value in row:
                    if value is None:
                        continue
                    for char in HAN_PATTERN.findall(str(value)):
                        if char not in seen:
                            seen.add(char)
                            chars.append(char)
    finally:
        workbook.close()

    return "".join(chars)


def main() -> None:
    parser = argparse.ArgumentParser(description="提取 Excel 中所有不重复汉字")
    parser.add_argument("input", nargs="?", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("-o", "--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()

    if not args.input.is_file():
        parser.error(f"Excel 文件不存在：{args.input}")

    chars = extract_unique_han_chars(args.input)
    args.output.parent.mkdir(parents=True, exist_ok=True)
    args.output.write_text(chars, encoding="utf-8")

    print(f"提取完成：{len(chars)} 个不同汉字")
    print(f"输出文件：{args.output.resolve()}")


if __name__ == "__main__":
    main()
